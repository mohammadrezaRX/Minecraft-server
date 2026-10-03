from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "word/انقلاب_دنیای_دیجیتال_نسخه_نمونه.docx"
FONT = "Arial"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(2.0)
    s.header_distance = Cm(0.8)
    s.footer_distance = Cm(0.8)

def rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")

def run(p, text, size=12, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:cs"), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)

def para(text="", size=12, bold=False, center=False):
    p = doc.add_paragraph()
    rtl(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if text:
        run(p, text, size, bold)
    return p

def heading(text, level=1):
    p = doc.add_paragraph()
    rtl(p)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run(p, text, 20 if level == 1 else 15, True, (31,71,120))
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; shade(c,"1F4778"); p=c.paragraphs[0]; rtl(p,WD_ALIGN_PARAGRAPH.CENTER); run(p,h,10,True,(255,255,255))
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            c=cells[j]
            if i%2==0: shade(c,"F5F8FB")
            p=c.paragraphs[0]; rtl(p); run(p,str(val),10)
    para("")

def flow(items):
    t=doc.add_table(rows=1, cols=len(items)*2-1)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,item in enumerate(items):
        c=t.cell(0,i*2); shade(c,["D9EAF7","E2F0D9","FFF2CC","FCE4D6","E4DFEC","DDEBF7"][i%6])
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; rtl(p,WD_ALIGN_PARAGRAPH.CENTER); run(p,item,9,True)
        if i<len(items)-1:
            p=t.cell(0,i*2+1).paragraphs[0]; rtl(p,WD_ALIGN_PARAGRAPH.CENTER); run(p,"←",18,True,(90,100,110))
    para("")

def box(title, body):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,"EEF5FB")
    p=c.paragraphs[0]; rtl(p); run(p,title,11,True,(31,71,120))
    p=c.add_paragraph(); rtl(p); run(p,body,10)
    para("")

# Header/footer
for sec in doc.sections:
    p=sec.header.paragraphs[0]; rtl(p); run(p,"انقلاب دنیای دیجیتال",9,True,(90,100,110))
    p=sec.footer.paragraphs[0]; rtl(p,WD_ALIGN_PARAGRAPH.CENTER)
    r=p.add_run(); fld=OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"),"begin"); instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "; end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); r._r.extend([fld,instr,end])

# Cover
para("پروژه پژوهشی فناوری و علوم رایانه",13,False,True)
para("انقلاب دنیای دیجیتال",31,True,True)
para("از تحول دیجیتال تا اینترنت اشیا و هوش مصنوعی",16,False,True)
doc.add_paragraph("")
flow(["انقلاب دیجیتال","شبکه‌ها","اینترنت اشیا","داده","هوش مصنوعی","سامانه هوشمند"])
for a,b in [("نام دانش‌آموز","................................"),("کلاس","................................"),("درس","................................"),("دبیر","................................")]:
    table(["مشخصات","مقدار"],[(a,b)])
doc.add_page_break()

# TOC
heading("فهرست مطالب")
table(["ردیف","عنوان","صفحه"],[
    ("۱","مقدمه","۳"),("۲","فصل اول: انقلاب دنیای دیجیتال","۴"),
    ("۳","فصل دوم: اینترنت اشیا (IoT)","۶"),("۴","فصل سوم: هوش مصنوعی (AI)","۸"),
    ("۵","ارتباط سه فناوری","۱۰"),("۶","نتیجه‌گیری","۱۱"),("۷","منابع","۱۲")])
doc.add_page_break()

heading("مقدمه")
para("فناوری دیجیتال در چند دهه اخیر شیوه ارتباط، آموزش، تولید، خرید، سرگرمی و دسترسی به اطلاعات را تغییر داده است. این تغییر گسترده را می‌توان با مفهوم انقلاب دنیای دیجیتال توضیح داد.")
para("در ادامه این تحول، دستگاه‌های بیشتری به شبکه‌ها متصل شدند و توانستند داده تولید و دریافت کنند؛ این ایده به اینترنت اشیا رسید. هم‌زمان، رشد حجم داده‌ها نیاز به روش‌های مؤثر برای تحلیل آن‌ها را افزایش داد و هوش مصنوعی به یکی از ابزارهای مهم این مرحله تبدیل شد.")
box("ایده اصلی پروژه","انقلاب دیجیتال بستر را گسترش می‌دهد → اینترنت اشیا داده تولید و منتقل می‌کند → هوش مصنوعی داده را تحلیل می‌کند → سامانه‌های هوشمند می‌توانند پیشنهاد، پیش‌بینی یا اقدام کنند.")
flow(["دیجیتالی شدن","شبکه‌ها","IoT","داده","AI","تصمیم / اقدام"])
doc.add_page_break()

heading("فصل اول: انقلاب دنیای دیجیتال")
para("انقلاب دیجیتال به مجموعه تغییراتی گفته می‌شود که در آن رایانه‌ها، فناوری‌های دیجیتال و شبکه‌های ارتباطی به‌طور گسترده وارد زندگی و فعالیت‌های انسانی شدند. این فرایند تدریجی بود و با پیشرفت رایانه، ارتباطات، ذخیره‌سازی اطلاعات و اینترنت شکل گرفت.")
flow(["اطلاعات فیزیکی","رایانه","شبکه دیجیتال","اینترنت","خدمات دیجیتال"])
heading("تأثیر در حوزه‌های مختلف",2)
table(["حوزه","نمونه تغییر"],[
    ("ارتباطات","پیام‌رسانی سریع، تماس تصویری و شبکه‌های اجتماعی"),
    ("آموزش","منابع آنلاین، کلاس مجازی و محتوای دیجیتال"),
    ("اقتصاد","فروشگاه اینترنتی، پرداخت و بانکداری دیجیتال"),
    ("صنعت","اتوماسیون و کنترل رایانه‌ای"),
    ("رسانه","پخش آنلاین و محتوای دیجیتال")])
heading("مزایا و چالش‌ها",2)
table(["مزایا","چالش‌ها"],[
    ("سرعت بیشتر انتقال اطلاعات","امنیت و حریم خصوصی"),
    ("دسترسی آسان‌تر به خدمات","وابستگی به زیرساخت دیجیتال"),
    ("خودکارسازی فرایندها","نیاز به مهارت‌های جدید"),
    ("کاهش برخی هزینه‌ها","مدیریت حجم زیاد داده")])
doc.add_page_break()

heading("فصل دوم: اینترنت اشیا (IoT)")
para("اینترنت اشیا به دستگاه‌ها و اشیای فیزیکی متصل به شبکه گفته می‌شود که با کمک حسگرها، نرم‌افزار و ارتباطات شبکه‌ای می‌توانند داده تولید، دریافت یا منتقل کنند.")
box("مثال ساده","یک حسگر دما می‌تواند دمای اتاق را اندازه‌گیری کند، داده را از طریق شبکه ارسال کند و یک سامانه بر اساس آن وضعیت سرمایش یا گرمایش را تنظیم کند.")
heading("چرخه کلی IoT",2)
flow(["حسگر","جمع‌آوری داده","شبکه","پردازش / ذخیره","تحلیل","اقدام / نمایش"])
heading("کاربردها",2)
table(["حوزه","مثال"],[
    ("خانه هوشمند","کنترل روشنایی، دما و برخی وسایل"),
    ("کشاورزی","اندازه‌گیری رطوبت خاک و شرایط محیط"),
    ("پزشکی","پایش برخی داده‌های سلامت با تجهیزات متصل"),
    ("صنعت","پایش دستگاه‌ها و نگهداری پیش‌بینانه"),
    ("حمل‌ونقل","ردیابی و مدیریت ناوگان"),
    ("شهر هوشمند","جمع‌آوری داده برای خدمات شهری")])
heading("چالش‌ها",2)
para("امنیت، حریم خصوصی، به‌روزرسانی نرم‌افزار، سازگاری دستگاه‌ها و مدیریت حجم داده‌ها از چالش‌های مهم اینترنت اشیا هستند.")
doc.add_page_break()

heading("فصل سوم: هوش مصنوعی (AI)")
para("هوش مصنوعی به روش‌ها و سامانه‌های رایانشی گفته می‌شود که برای انجام وظایفی مانند تشخیص الگو، پیش‌بینی، طبقه‌بندی، پردازش زبان و استنتاج از داده طراحی می‌شوند. یادگیری ماشین یکی از بخش‌های مهم هوش مصنوعی است.")
flow(["داده / ورودی","آماده‌سازی","مدل AI","یادگیری / استنتاج","خروجی"])
heading("کاربردها",2)
table(["حوزه","مثال"],[
    ("پزشکی","کمک به تحلیل تصاویر پزشکی"),
    ("آموزش","شخصی‌سازی برخی محتواها و تحلیل عملکرد"),
    ("زبان","ترجمه، تشخیص گفتار و تولید متن"),
    ("بینایی ماشین","تشخیص اشیا و الگوهای تصویری"),
    ("صنعت","پیش‌بینی خرابی و تحلیل فرایند"),
    ("حمل‌ونقل","تحلیل ترافیک و سامانه‌های هوشمند")])
box("نکته","کیفیت داده، روش طراحی مدل و نحوه ارزیابی آن بر کیفیت خروجی هوش مصنوعی اثر می‌گذارد.")
doc.add_page_break()

heading("ارتباط سه فناوری")
para("این سه مفهوم را می‌توان بخش‌های مرتبط یک تحول فناوری دانست. انقلاب دیجیتال بستر دیجیتال و شبکه‌ای را گسترش داد؛ اینترنت اشیا ارتباط دنیای فیزیکی و دیجیتال را بیشتر کرد؛ و هوش مصنوعی امکان تحلیل داده‌ها و استخراج الگوها را افزایش داد.")
heading("مثال: خانه هوشمند",2)
flow(["حسگر دما","اندازه‌گیری","ارسال داده","ذخیره / پردازش","تحلیل با AI","پیشنهاد / اقدام"])
para("در این نمونه، حسگر بخشی از اینترنت اشیاست، شبکه داده را جابه‌جا می‌کند و هوش مصنوعی می‌تواند برای تحلیل الگوهای مصرف یا شرایط محیط استفاده شود.")
heading("مقایسه",2)
table(["موضوع","نقش اصلی","مثال"],[
    ("انقلاب دیجیتال","ایجاد و گسترش بستر دیجیتال","خدمات دیجیتال"),
    ("اینترنت اشیا","تولید و تبادل داده از اشیای فیزیکی","خانه هوشمند"),
    ("هوش مصنوعی","تحلیل، پیش‌بینی و استنتاج","تشخیص تصویر")])
box("رابطه کلی","انقلاب دیجیتال → شبکه‌های گسترده → اینترنت اشیا → داده → هوش مصنوعی → سامانه‌های هوشمند")
doc.add_page_break()

heading("نتیجه‌گیری")
para("دنیای دیجیتال نتیجه یک فناوری منفرد نیست، بلکه حاصل کنار هم قرار گرفتن رایانه‌ها، شبکه‌ها، نرم‌افزارها، داده‌ها و روش‌های هوشمند است. با گسترش فناوری دیجیتال، دستگاه‌های بیشتری به شبکه متصل شدند و مقدار بیشتری داده تولید شد. روش‌های هوش مصنوعی می‌توانند برای تحلیل این داده‌ها و استخراج الگوها به کار گرفته شوند.")
heading("نمودار نهایی",2)
flow(["انقلاب دیجیتال","اینترنت و شبکه","IoT","داده","AI","سامانه هوشمند","آینده دیجیتال"])
box("جمع‌بندی","سه موضوع این پروژه به هم وابسته‌اند: انقلاب دیجیتال بستر را فراهم می‌کند، اینترنت اشیا داده را از دنیای واقعی وارد سامانه‌ها می‌کند و هوش مصنوعی می‌تواند از داده برای تحلیل و استنتاج استفاده کند.")
doc.add_page_break()

heading("منابع")
for s in [
    "IBM — مطالب آموزشی هوش مصنوعی و یادگیری ماشین",
    "Oracle — مطالب آموزشی اینترنت اشیا (IoT)",
    "Microsoft Learn — آموزش‌های هوش مصنوعی و فناوری دیجیتال",
    "Encyclopaedia Britannica — مطالب مرتبط با انقلاب دیجیتال و فناوری اطلاعات",
    "International Telecommunication Union (ITU) — مطالب و گزارش‌های فناوری اطلاعات و ارتباطات"]:
    p=doc.add_paragraph(); rtl(p); run(p,"• "+s,11)

doc.core_properties.title="انقلاب دنیای دیجیتال"
doc.core_properties.subject="انقلاب دیجیتال، اینترنت اشیا و هوش مصنوعی"
doc.core_properties.author="پروژه دانش‌آموزی"
doc.save(OUT)
print(OUT)
